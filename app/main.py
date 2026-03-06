from __future__ import annotations

import json
import random
import re
import uuid
from pathlib import Path
from typing import Any

from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
UPLOAD_DIR = DATA_DIR / "uploads"
BANK_PATH = DATA_DIR / "question_bank.json"

DATA_DIR.mkdir(parents=True, exist_ok=True)
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

app = FastAPI(title="Smart Exam Builder")
app.mount("/static", StaticFiles(directory=BASE_DIR / "static"), name="static")

EXAM_SESSIONS: dict[str, dict[str, Any]] = {}


@app.get("/")
def root() -> FileResponse:
    return FileResponse(BASE_DIR / "static" / "index.html")


def load_bank() -> list[dict[str, Any]]:
    if not BANK_PATH.exists():
        return []
    try:
        return json.loads(BANK_PATH.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return []


def save_bank(questions: list[dict[str, Any]]) -> None:
    BANK_PATH.write_text(json.dumps(questions, ensure_ascii=False, indent=2), encoding="utf-8")


def read_upload_text(file_path: Path) -> str:
    suffix = file_path.suffix.lower()

    if suffix == ".txt":
        return file_path.read_text(encoding="utf-8", errors="ignore")

    if suffix == ".pdf":
        try:
            import pdfplumber  # type: ignore
        except ImportError as exc:
            raise HTTPException(status_code=500, detail="缺少 pdfplumber，請先安裝 requirements") from exc

        parts: list[str] = []
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                parts.append(page.extract_text() or "")
        return "\n".join(parts)

    if suffix == ".docx":
        try:
            from docx import Document  # type: ignore
        except ImportError as exc:
            raise HTTPException(status_code=500, detail="缺少 python-docx，請先安裝 requirements") from exc

        doc = Document(file_path)
        return "\n".join(p.text for p in doc.paragraphs)

    if suffix in {".xlsx", ".xls"}:
        try:
            import openpyxl  # type: ignore
        except ImportError as exc:
            raise HTTPException(status_code=500, detail="缺少 openpyxl，請先安裝 requirements") from exc

        wb = openpyxl.load_workbook(file_path, data_only=True)
        rows: list[str] = []
        for ws in wb.worksheets:
            rows.append(f"### 工作表: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if vals:
                    rows.append(" | ".join(vals))
        return "\n".join(rows)

    raise HTTPException(status_code=400, detail=f"不支援的檔案格式: {suffix}")


def normalize_option_key(raw: str) -> str:
    raw = raw.strip().upper()
    m = re.search(r"[A-F]", raw)
    return m.group(0) if m else raw[:1]


def parse_question_blocks(raw_text: str) -> list[dict[str, Any]]:
    # Normalize line endings and remove obvious page artifacts.
    text = raw_text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"\n{3,}", "\n\n", text)

    blocks = [b.strip() for b in re.split(r"\n\s*\n", text) if b.strip()]
    questions: list[dict[str, Any]] = []

    for block in blocks:
        lines = [ln.strip() for ln in block.split("\n") if ln.strip()]
        if len(lines) < 2:
            continue

        q_text_parts: list[str] = []
        options: list[dict[str, str]] = []
        answer = ""
        explanation = ""

        current_opt_idx: int | None = None

        for line in lines:
            ans_match = re.match(r"^(?:答案|answer)\s*[:：]\s*([A-Fa-f])", line, flags=re.IGNORECASE)
            exp_match = re.match(r"^(?:解析|explanation)\s*[:：]\s*(.+)", line, flags=re.IGNORECASE)
            opt_match = re.match(r"^([A-Fa-f])[\)\]\.、:：\s-]+(.+)$", line)

            if ans_match:
                answer = normalize_option_key(ans_match.group(1))
                current_opt_idx = None
                continue

            if exp_match:
                explanation = exp_match.group(1).strip()
                current_opt_idx = None
                continue

            if opt_match:
                key = normalize_option_key(opt_match.group(1))
                options.append({"key": key, "text": opt_match.group(2).strip()})
                current_opt_idx = len(options) - 1
                continue

            if current_opt_idx is not None:
                options[current_opt_idx]["text"] += f" {line}"
            else:
                clean_line = re.sub(r"^(?:\d+|Q\d+|題\s*\d+)\s*[\)\]\.、:：-]?\s*", "", line, flags=re.IGNORECASE)
                q_text_parts.append(clean_line)

        if len(options) >= 2 and q_text_parts:
            q_text = " ".join(q_text_parts).strip()
            if not q_text:
                continue
            if not answer:
                # Skip question without answer for MVP reliability.
                continue

            questions.append(
                {
                    "id": str(uuid.uuid4()),
                    "question": q_text,
                    "options": options,
                    "answer": answer,
                    "explanation": explanation,
                }
            )

    return questions


@app.post("/api/upload")
async def upload_file(file: UploadFile = File(...)) -> dict[str, Any]:
    if not file.filename:
        raise HTTPException(status_code=400, detail="缺少檔案名稱")

    ext = Path(file.filename).suffix.lower()
    if ext not in {".pdf", ".docx", ".xlsx", ".xls", ".txt"}:
        raise HTTPException(status_code=400, detail="只支援 PDF / DOCX / XLSX / XLS / TXT")

    save_path = UPLOAD_DIR / f"{uuid.uuid4()}{ext}"
    data = await file.read()
    save_path.write_bytes(data)

    text = read_upload_text(save_path)
    parsed = parse_question_blocks(text)
    if not parsed:
        raise HTTPException(
            status_code=400,
            detail="沒有抽出有效題目。請確認題目格式含題幹、A/B/C/D選項與答案欄位。",
        )

    current = load_bank()
    current.extend(parsed)
    save_bank(current)

    return {
        "message": "上傳成功並完成題目抽取",
        "added_count": len(parsed),
        "total_count": len(current),
    }


@app.get("/api/questions")
def list_questions() -> dict[str, Any]:
    qs = load_bank()
    masked = [
        {
            "id": q["id"],
            "question": q["question"],
            "options": q["options"],
        }
        for q in qs
    ]
    return {"count": len(masked), "questions": masked}


@app.post("/api/exam/start")
def start_exam(payload: dict[str, Any]) -> dict[str, Any]:
    count = int(payload.get("count", 10))
    bank = load_bank()

    if not bank:
        raise HTTPException(status_code=400, detail="題庫是空的，請先上傳題目")

    count = max(1, min(count, len(bank)))
    selected = random.sample(bank, count)
    exam_id = str(uuid.uuid4())

    EXAM_SESSIONS[exam_id] = {"question_ids": [q["id"] for q in selected]}

    public_q = [
        {
            "id": q["id"],
            "question": q["question"],
            "options": q["options"],
        }
        for q in selected
    ]
    return {"exam_id": exam_id, "count": len(public_q), "questions": public_q}


@app.post("/api/exam/submit")
def submit_exam(payload: dict[str, Any]) -> dict[str, Any]:
    exam_id = payload.get("exam_id", "")
    answers: dict[str, str] = payload.get("answers", {})

    if exam_id not in EXAM_SESSIONS:
        raise HTTPException(status_code=404, detail="找不到考試場次，請重新開始")

    session = EXAM_SESSIONS[exam_id]
    qid_set = set(session["question_ids"])

    bank = load_bank()
    q_map = {q["id"]: q for q in bank if q["id"] in qid_set}

    total = len(q_map)
    correct = 0
    details: list[dict[str, Any]] = []

    for qid, q in q_map.items():
        user_ans = normalize_option_key(answers.get(qid, "")) if answers.get(qid) else ""
        real_ans = normalize_option_key(q.get("answer", ""))
        is_correct = user_ans == real_ans and bool(user_ans)
        if is_correct:
            correct += 1

        details.append(
            {
                "id": qid,
                "question": q["question"],
                "your_answer": user_ans,
                "correct_answer": real_ans,
                "is_correct": is_correct,
                "explanation": q.get("explanation", ""),
            }
        )

    score = round(correct / total * 100, 2) if total else 0
    return {"score": score, "correct": correct, "total": total, "details": details}


@app.post("/api/questions/clear")
def clear_questions() -> dict[str, Any]:
    save_bank([])
    EXAM_SESSIONS.clear()
    return {"message": "題庫已清空"}
